"""Fuente editorial reproducible del curso de técnicos de exterior.

El constructor del curso reutiliza estas reglas para mantener alineados el
manual, las diapositivas y las transcripciones de los 92 audios.
"""

from __future__ import annotations

import re
from pathlib import Path
from collections import Counter

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SLIDES_TXT = ROOT / "content/tecnicos-exterior-20h/slides.txt"
MANUAL_TXT = ROOT / "content/tecnicos-exterior-20h/manual.txt"
LOGO = ROOT / "public/images/inminer-campus-logo.png"
OUT_DIR = ROOT / "tmp/tecnicos-exterior-20h"
OUT = OUT_DIR / "Guion_Maestro_ElevenLabs_Tecnicos_Exterior_20h_v1.1_Audios_2_3_min.docx"


# Cada registro aporta la intención docente, un ejemplo de campo y el mensaje que debe quedar.
# El texto final se completa con el contenido de la diapositiva y del manual del curso.
META_RAW = r"""
1|Abrir el curso situando al técnico como la persona que conecta planificación, producción y prevención.|Imagine el inicio de un turno tras una noche de lluvia. La producción quiere arrancar, pero antes alguien debe comprobar pistas, frentes, drenajes, máquinas y coordinación. Esa mirada global es la función técnica.|Anticipar, supervisar y poder parar son tres partes de una misma responsabilidad.
2|Distinguir peligro, exposición, riesgo, control, verificación y evidencia sin tratarlos como sinónimos.|Una roca inestable es un peligro. Un equipo que trabaja bajo su trayectoria crea exposición. La combinación determina el riesgo. La exclusión controla la exposición y una inspección registrada demuestra que la barrera existe.|Un control no está implantado hasta que puede verificarse y dejar evidencia.
3|Explicar para qué sirve cada documento de la explotación y evitar usar uno como sustituto de otro.|Un procedimiento de mantenimiento puede ser impecable y, aun así, resultar inseguro si contradice la geometría aprobada, el Documento sobre Seguridad y Salud o una disposición interna de seguridad.|Cada documento responde a una pregunta distinta y debe ser coherente con los de nivel superior.
4|Diferenciar procedimiento, instrucción y permiso de trabajo por su función práctica.|Para intervenir una trituradora puede existir un procedimiento general, una instrucción de consignación paso a paso y un permiso que autorice esa intervención concreta, con responsables, hora y condiciones.|El documento correcto debe llegar a la persona correcta antes de iniciar la tarea.
5|Aplicar la jerarquía de controles priorizando diseño y eliminación de exposición.|Si peatones y volquetes comparten una zona, repartir chalecos no elimina la interacción. Rediseñar el recorrido, separar físicamente o suprimir el cruce reduce la dependencia de la conducta.|Cuanto más depende una barrera de recordar y obedecer, menos robusta es.
6|Activar la gestión del cambio antes de continuar cuando varían terreno, agua, equipo, método u organización.|Una motoniveladora cambia una pendiente transversal; una lluvia revela que el agua ya no evacua; aparece barro en una curva y cambia la trayectoria de los camiones. El cambio aparentemente pequeño ha modificado el riesgo.|Cambiar primero y evaluar después invierte el orden preventivo.
7|Guiar la actividad sobre un cambio pequeño que altera drenaje y tráfico.|El grupo debe identificar qué cambió, quién queda expuesto, qué señales mostrarían deterioro y qué acción inmediata permitiría controlar la situación sin improvisar.|La actividad se resuelve con decisiones observables, no con frases genéricas.
8|Explicar la solución del caso mostrando control inmediato, corrección de causa y verificación de eficacia.|Cerrar o regular el tráfico evita la exposición inmediata; restablecer el drenaje corrige la causa; observar el comportamiento tras nuevas lluvias verifica que la solución funciona.|Cerrar una incidencia exige comprobar el resultado, no solo ejecutar una reparación.
9|Comprobar conceptos fundamentales del marco de partida antes de avanzar.|El alumno debe comparar eliminación, señalización, EPI e instrucciones; distinguir Documento sobre Seguridad y Salud y disposición interna; y reconocer cuándo un cambio obliga a reevaluar.|La comprobación debe revelar razonamiento, no memoria literal.
10|Razonar las respuestas del bloque cero y corregir errores frecuentes.|Una persona puede conocer una norma y seguir expuesta si el diseño conserva la trayectoria del peligro. Del mismo modo, un papel firmado no demuestra que una barrera esté presente en campo.|La mejor respuesta siempre conecta peligro, exposición, control y evidencia.
11|Presentar el ciclo completo de una explotación exterior como una cadena de tareas relacionadas.|Desde la preparación y el arranque hasta el tratamiento, mantenimiento y restauración, cada fase entrega condiciones a la siguiente. Un defecto de una etapa puede reaparecer lejos y horas después.|Evaluar bien exige mirar el proceso completo, no una tarea aislada.
12|Identificar interacciones críticas entre preparación, arranque, carga, transporte y tratamiento.|Una voladura afecta al saneo; el saneo condiciona la carga; la carga influye en estabilidad y tráfico; el transporte modifica pistas y el tratamiento puede devolver atascos o polvo al resto del sistema.|Las interfaces entre fases merecen tanta atención como cada fase por separado.
13|Enseñar a describir una tarea con un mínimo común de datos antes de evaluar el riesgo.|No basta escribir mantenimiento de cinta. Hay que precisar tramo, aislamiento, personal, herramientas, secuencia, duración, energías, trabajos simultáneos, criterios de parada y estado de entrega.|Una definición incompleta produce una evaluación genérica y controles débiles.
14|Ordenar el análisis alrededor de personas, equipos y energías presentes.|En el cambio de un latiguillo intervienen mecánicos, operador, supervisión, excavadora, bomba hidráulica, implemento, presión residual y gravedad. Olvidar uno solo puede abrir una trayectoria peligrosa.|Las personas no se analizan separadas de los equipos y las energías que las rodean.
15|Aplicar el modelo energía, trayectoria y persona expuesta a cualquier tarea.|La energía de una carga suspendida sigue una trayectoria vertical; la de una manguera presurizada puede proyectar fluido; la de un volquete en pendiente sigue la dirección de movimiento posible.|Si no se puede describir la trayectoria, todavía no se entiende el escenario.
16|Reconocer interferencias incluso cuando cada tarea dispone de su propio procedimiento.|Producción quiere probar un triturador mientras una contrata trabaja aguas abajo. Ambas tareas pueden ser correctas por separado, pero su coincidencia introduce energía en una zona ocupada.|Coordinar significa controlar la interacción real, no intercambiar documentos.
17|Definir criterios de parada, evidencias y autoridad de reanudación antes de empezar.|Una grieta nueva, pérdida de visibilidad, fuga presurizada o señal de barreno fallido deben tener una respuesta acordada. También debe saberse quién puede confirmar que el riesgo ha sido reevaluado.|Parar es eficaz solo si existe un camino claro y documentado para reanudar.
18|Dirigir una actividad en la que el alumno describa una operación real antes de evaluarla.|Puede elegirse carga de volquetes, limpieza de tolva o mantenimiento de cinta. El grupo debe reconstruir pasos, ubicaciones, participantes, energías, interferencias, parada y entrega final.|Primero se describe el trabajo real; después se valora su riesgo.
19|Comprobar el orden del ciclo y la calidad mínima de una ficha de tarea.|El alumno debe detectar que una ficha sin criterios de parada ni autoridad de reanudación no permite controlar desviaciones, aunque enumere equipos y personas.|Una ficha útil permite decidir antes, durante y al terminar.
20|Explicar las respuestas del bloque uno conectándolas con el trabajo de campo.|Ordenar fases no es un ejercicio académico: permite ver qué condiciones recibe cada equipo. Identificar energía y trayectoria convierte una lista en un escenario controlable.|La definición del trabajo es la base de todo el análisis posterior.
21|Presentar un método operativo de evaluación que no se reduzca a una matriz numérica.|Se describe el peligro, la trayectoria y la exposición; después se estima consecuencia y probabilidad, se eligen controles, se asigna responsable y se define cómo verificar eficacia.|Una cifra ayuda a priorizar, pero nunca sustituye la descripción del escenario.
22|Explicar los factores que gobiernan la estabilidad de taludes, bancos, bermas y plataformas.|Dos frentes con el mismo ángulo pueden comportarse de forma distinta por discontinuidades, agua, meteorización, orientación o daños de voladura. El proyecto fija la geometría, pero el campo confirma su estado.|La estabilidad se gestiona combinando diseño, observación y revisión ante cambios.
23|Reconocer indicadores tempranos de inestabilidad y responder sin acercarse al peligro.|Grietas nuevas, caída de fragmentos, abombamiento, agua turbia o deformación de bermas son información. La respuesta inicial es retirar exposición, delimitar, comunicar y evaluar desde un lugar seguro.|No se entra en una zona potencialmente inestable para comprobar si realmente lo está.
24|Diferenciar inspección operativa, técnica y especializada y saber cuándo escalar.|El operador puede detectar una grieta; una inspección técnica compara evolución y criterios; un especialista geotécnico interpreta mecanismos y define instrumentación o rediseño.|Cada nivel observa algo distinto y ninguno debe simular competencias que no posee.
25|Relacionar agua, drenaje y meteorología con estabilidad, pistas y capacidad de frenado.|El agua puede saturar un material, erosionar un pie, ocultar un bache y reducir adherencia. Un bombeo improvisado también puede trasladar el problema a otra ladera o escombrera.|El agua se controla desde su origen hasta un punto de vertido seguro.
26|Plantear el caso de lluvia posterior a una voladura y obligar a decidir antes de entrar.|El alumno debe tratar conjuntamente agua turbia, material suelto y grieta nueva, establecer exclusión, comunicar, pedir evaluación competente y definir condiciones de reapertura.|La ausencia de un desprendimiento previo no es evidencia de estabilidad.
27|Explicar que la geometría de las pistas es una barrera técnica y debe comprobarse contra proyecto e ITC.|Anchura, arcenes, distancia al borde, pendiente, radios, visibilidad, drenaje y pretiles forman un sistema. Una cifra aislada no garantiza seguridad si el resto del diseño falla.|La pista debe ser compatible con el vehículo real y con las condiciones reales.
28|Controlar la interacción entre maquinaria móvil y personas mediante separación y reglas visibles.|Un topógrafo que entra en el radio de una cargadora no queda protegido solo por un chaleco. Se necesita exclusión, acceso definido, contacto con el operador y ubicación fuera de puntos ciegos.|La prioridad es evitar que persona y máquina compartan trayectoria.
29|Enseñar el contacto positivo antes de entrar en una zona de maniobra o iniciar marcha atrás.|La persona solicita acceso desde un lugar visible, el operador detiene y confirma, y solo entonces se autoriza la aproximación. Si se pierde contacto, la maniobra se detiene.|Ver la máquina no significa que el operador haya visto a la persona.
30|Explicar los puntos ciegos como limitación de diseño que exige controles redundantes.|Cámaras, espejos y alarmas ayudan, pero pueden ensuciarse, fallar o no cubrir todo el entorno. El diseño del flujo, las zonas de espera y la eliminación de marcha atrás son más robustos.|Un avisador no convierte una zona ciega en una zona segura.
31|Delimitar qué coordina el técnico alrededor de perforación y voladura sin atribuirle habilitaciones inexistentes.|El técnico verifica terreno, estabilidad, accesos, interferencias, horario, meteorología, evacuación y recuento. La manipulación de explosivos corresponde al personal habilitado.|Coordinar el entorno de una voladura no convierte a nadie en artillero.
32|Fijar una respuesta inequívoca ante un barreno sospechoso o fallido.|Si aparece una perforación, cable, detonador o zona dudosa, se detiene la excavación, se mantiene distancia, se señaliza y se avisa por el canal previsto. Nadie investiga manipulando.|La duda se gestiona como barreno fallido hasta que personal autorizado determine lo contrario.
33|Mostrar todas las formas de energía peligrosa que pueden permanecer tras detener un equipo.|Además de electricidad, existen presión hidráulica y neumática, gravedad, tensión mecánica, calor, rotación, material almacenado y posibilidad de reenergización remota.|Poner un mando a cero no equivale a aislar.
34|Explicar una consignación completa en nueve pasos y la responsabilidad individual sobre el bloqueo.|Preparar, parar, aislar, bloquear e identificar, disipar, verificar, ejecutar, probar de forma controlada y retirar bloqueos son una secuencia. Saltar la verificación deja una suposición en el centro del trabajo.|El estado de energía cero debe demostrarse antes de tocar.
35|Dirigir el caso del atasco en una cinta con tensión, material y rearme remoto.|El alumno debe identificar energía eléctrica, tensión de banda, gravedad del material y acción de terceros; proponer aislamiento aguas arriba y abajo y definir una prueba de ausencia de energía.|El atasco se libera solo cuando todas las fuentes están controladas.
36|Explicar por qué el polvo visible no representa toda la exposición a sílice cristalina respirable.|Las partículas respirables pueden permanecer suspendidas y penetrar profundamente sin formar una nube espectacular. Perforación, trituración, transferencia y limpieza son focos principales.|No se evalúa la exposición mirando únicamente el aire.
37|Aplicar controles de polvo en fuente, trayectoria, organización y persona.|Método húmedo, captación, cerramiento y automatización actúan en origen; cabinas filtradas y separación controlan trayectoria; mantenimiento y limpieza sostienen la eficacia; el respirador queda como complemento.|Una mascarilla no compensa una captación averiada o una cabina con fugas.
38|Integrar ruido, vibraciones y calor en la planificación diaria y no tratarlos como molestias menores.|Un operador puede acumular ruido, vibración de cuerpo entero y carga térmica durante el mismo turno. Mantenimiento, asiento, velocidad, pausas, sombra e hidratación deben coordinarse.|La exposición total depende de intensidad, duración y condiciones de la persona.
39|Desarrollar electricidad, líneas aéreas, incendio y sustancias con criterios claros de parada.|Una caja elevada puede invadir la distancia de seguridad de una línea; una fuga de combustible puede alcanzar una fuente caliente. Para trabajos sin tensión deben cumplirse las cinco reglas de oro completas.|La energía eléctrica y el fuego exigen distancias, aislamiento y personal competente.
40|Explicar la respuesta de emergencia con la secuencia proteger, avisar, socorrer, recibir y preservar.|Antes de acercarse a una víctima se comprueban tráfico, rocas, electricidad, fuego y atmósfera. Después se activa el plan con ubicación exacta y se guía a los servicios externos.|Ayudar nunca debe crear una segunda víctima.
41|Comprobar consignación, puntos ciegos, polvo y emergencias mediante preguntas razonadas.|Ordenar aislar, bloquear, disipar y verificar revela si el alumno entiende el estado seguro. Negar que una alarma elimine un punto ciego evita confiar en una sola ayuda.|La comprobación debe centrarse en decisiones críticas.
42|Resolver los casos del bloque dos mostrando por qué cada control es necesario.|La lluvia exige exclusión y evaluación competente; la cinta exige consignación integral; el polvo exige controlar la fuente y verificar la exposición.|La solución correcta elimina o controla la trayectoria antes de recuperar producción.
43|Presentar el ciclo de vida del equipo desde selección hasta retirada.|Un equipo puede ser adecuado al comprarlo y dejar de serlo tras una modificación, cambio de implemento, envejecimiento o nuevo terreno. Cada fase necesita criterios y registros.|La conformidad inicial no sustituye la gestión durante toda la vida del equipo.
44|Comparar riesgos dominantes de perforadora, excavadora y cargadora sin caer en listas genéricas.|La perforadora combina estabilidad, polvo y elementos rotativos; la excavadora añade alcance y frente; la cargadora incorpora articulación, atropello y caída de material.|La evaluación debe seguir el movimiento y la función real de cada máquina.
45|Comparar volquete, bulldócer y motoniveladora atendiendo a geometría, visibilidad y tarea.|El volquete depende de frenos, borde y caja; el bulldócer trabaja cerca de desniveles y con visibilidad limitada; la motoniveladora ocupa gran anchura y altera drenajes.|La misma pista no ofrece el mismo riesgo a todas las máquinas.
46|Explicar riesgos de trituración, cintas y vehículos ligeros y sus interacciones.|Una cinta puede arrancar desde otro punto; una trituradora conserva material y energía; un vehículo ligero puede desaparecer en el punto ciego de un volquete.|Los equipos fijos y móviles forman un sistema conectado.
47|Aclarar qué demuestran documentación, declaración de conformidad y marcado CE y qué no demuestran.|La identificación y el manual permiten saber configuración, límites y mantenimiento. El marcado CE no certifica que el terreno, la ruta o la tarea concreta sean seguros.|La adecuación al puesto se verifica en la explotación.
48|Explicar la función conjunta de estructuras ROPS y FOPS y del cinturón de seguridad.|En un vuelco, la estructura crea un volumen de supervivencia y el cinturón mantiene al operador dentro. Una estructura dañada o un cinturón defectuoso requieren evaluación e inmovilización según criticidad.|ROPS sin cinturón no forma un sistema completo de protección.
49|Revisar frenos, dirección, visibilidad y avisadores como funciones críticas antes del uso.|Una prueba debe hacerse en condiciones controladas y comparar el resultado con el criterio del fabricante. La ausencia de testigos no demuestra capacidad de frenado.|Las funciones críticas se comprueban, no se presuponen.
50|Diferenciar guardas y paradas de emergencia y evitar usar la parada como medio habitual de aislamiento.|La guarda impide contacto; el enclavamiento controla acceso; la parada de emergencia limita consecuencias. Para intervenir se requiere consignación, porque el equipo puede rearmarse.|Una seta de emergencia no sustituye un bloqueo personal.
51|Enseñar un recorrido sistemático de inspección antes del uso.|Se comienza por el entorno, se recorre estructura, rodaje, fluidos, cabina, mandos, protecciones y función. El equipo permanece inmovilizado y nadie entra bajo elementos elevados.|La secuencia estable reduce olvidos y mejora la calidad del registro.
52|Clasificar defectos como críticos, limitantes o menores con criterio conservador.|Un fallo de freno o cinturón defectuoso en un equipo con riesgo de vuelco es crítico. Un daño cosmético sin función puede ser menor, pero solo después de confirmar que no afecta a seguridad.|Si no se conoce la consecuencia, el defecto no se rebaja por intuición.
53|Dirigir una inspección de campo basada en una máquina real o fotografías.|El alumno debe identificar equipo, seguir el recorrido, registrar defectos, clasificarlos, decidir acción y justificar qué evidencia permitiría cerrar cada uno.|Inspeccionar no es marcar casillas; es producir una decisión segura.
54|Explicar mantenimiento seguro y fluidos a presión con planificación previa.|Se elige lugar, acceso, iluminación, herramientas, soportes, repuestos y control de derrames. Las fugas no se buscan con la mano y la presión se libera según fabricante.|La intervención empieza antes de abrir una protección o aflojar una conexión.
55|Controlar implementos elevados, trabajos en caliente y riesgos de caída o incendio.|Un implemento elevado se baja o soporta mecánicamente; un tablón no es soporte calculado. El trabajo en caliente necesita limpieza, separación de combustibles, extintor y vigilancia posterior.|La hidráulica y la ausencia de llama visible no son garantías suficientes.
56|Situar el EPI como último nivel de control y explicar selección, compatibilidad y mantenimiento.|Casco, gafas, protección auditiva y respiratoria pueden interferir entre sí. Deben ajustarse a la persona y al riesgo residual, revisarse y sustituirse cuando pierdan eficacia.|Entregar un EPI no demuestra que proteja durante la tarea.
57|Relacionar riesgos concretos con selección y ajuste del EPI.|Polvo exige filtro y ajuste facial; ruido exige atenuación adecuada sin impedir comunicaciones críticas; altura exige sistema compatible con punto de anclaje y distancia de caída.|El EPI se selecciona desde la evaluación, no desde un catálogo genérico.
58|Diferenciar escalera, andamio y plataforma elevadora y matizar cuándo una escalera puede usarse como puesto.|La escalera se reserva para casos justificados de bajo riesgo y corta duración cuando no sea razonable otro medio. Andamios y plataformas requieren montaje, inspección, nivelación y protección colectiva.|El medio de acceso debe elegirse por riesgo y duración, no por comodidad.
59|Planificar izados considerando carga, equipo, accesorios, terreno y comunicaciones.|Antes de elevar se conoce masa, centro de gravedad, configuración, capacidad, radio, suelo, señalista y zona de exclusión. Nadie pasa bajo la carga.|Un accesorio certificado puede fallar si se usa con un ángulo o configuración no previstos.
60|Resolver los casos de cámara averiada y cambio de latiguillo sin compensaciones improvisadas.|Una alarma no reemplaza visibilidad crítica; el volquete se inmoviliza o se rediseña la maniobra bajo evaluación. El latiguillo exige bajar, soportar, despresurizar, verificar y controlar derrames.|La producción no convierte un defecto crítico en aceptable.
61|Comprobar clasificación de defectos, marcado CE, EPI y medios auxiliares.|Las preguntas deben obligar a justificar por qué un freno, un cinturón o una cámara afectan a funciones de seguridad y por qué la conformidad no cubre cualquier terreno.|Clasificar es relacionar defecto, función y consecuencia.
62|Razonar las respuestas del bloque tres y reforzar decisiones conservadoras.|La cámara averiada cambia la maniobra; el latiguillo conserva presión; el marcado CE no evalúa la explotación; la escalera no es siempre la opción correcta.|La evidencia de reparación y prueba funcional cierra el defecto.
63|Definir los objetivos de la vigilancia y los elementos que deben observarse.|Vigilar significa comparar con un estándar, detectar cambios, actuar y verificar cierre. Incluye frentes, bermas, escombreras, pistas, aguas, instalaciones, equipos y conductas.|Una ronda sin criterio de aceptación produce observaciones, no control.
64|Diseñar frecuencias de inspección y disparadores extraordinarios según riesgo y cambio.|La lluvia, una voladura, un sismo, una modificación, un fallo repetido o una ausencia prolongada pueden exigir inspección adicional, aunque la ronda prevista ya se haya realizado.|La frecuencia fija nunca sustituye la vigilancia provocada por eventos.
65|Explicar una inspección que produce decisiones desde la preparación hasta la verificación.|El técnico prepara estándar, observa desde zona segura, recorre una secuencia, contrasta con ejecutantes, clasifica, asigna acciones y vuelve a comprobar.|El cierre ocurre en campo cuando la función se ha recuperado.
66|Registrar desviaciones con responsable, plazo, criticidad y criterio de cierre.|Una fotografía puede mostrar un bache relleno, pero no demostrar drenaje ni comportamiento bajo carga. El registro debe decir qué resultado se espera y cómo se verificará.|Sin criterio de cierre, una acción puede quedar terminada administrativamente y abierta técnicamente.
67|Diferenciar estados normal, alerta y alarma y asociar una respuesta previamente definida.|Una tendencia de deformación puede activar alerta y mayor vigilancia; superar un umbral o perder una barrera crítica activa alarma, parada y exclusión.|Los umbrales deben acordarse antes de que la presión operativa influya en la decisión.
68|Convertir el derecho y deber de parada en una capacidad operativa protegida.|Quien detecta una condición peligrosa debe saber cómo avisar, dejar el equipo seguro, controlar la zona y recibir respuesta. La organización evita represalias y exige cierre.|Parar a tiempo es una acción profesional, no un fracaso productivo.
69|Explicar la reapertura documentada después de una parada.|Se identifica causa, se implantan controles, se comprueba función, se comunica y una persona autorizada decide reanudar. Que la señal desaparezca por sí sola no demuestra seguridad.|La reapertura necesita evidencia proporcional a la criticidad.
70|Resolver el caso del bache repetido buscando la causa sistémica.|Rellenar tres veces sin analizar drenaje, base, pendiente, carga o método demuestra repetición, no control. Primero se regula tráfico; después se corrige la causa y se verifica durante un periodo.|Una desviación recurrente debe escalarse y analizarse como tendencia.
71|Realizar una comprobación con opciones completas sobre alarma y orden de inspección.|La opción correcta para alarma debe incluir parar, aislar, excluir, evaluar y autorizar reapertura. La secuencia de inspección comienza preparando el estándar y termina verificando el cierre.|La respuesta correcta conserva el orden preventivo.
72|Explicar la resolución del caso y las respuestas del bloque cuatro.|Controlar el tráfico reduce exposición inmediata; analizar drenaje y estructura trata la causa; observar el comportamiento posterior demuestra eficacia.|Repetir una reparación fallida no equivale a cerrar la desviación.
73|Definir interferencia y riesgo emergente entre tareas, empresas o fases.|Una empresa puede cumplir su procedimiento y, aun así, energizar el equipo de otra. La interferencia aparece en espacio, tiempo, energía, información o emergencia.|La suma de dos trabajos seguros por separado puede ser insegura.
74|Analizar las dimensiones espacio, tiempo, energía, información, acceso y emergencia.|Un equipo puede compartir zona, una prueba puede coincidir con mantenimiento, una alarma puede no oírse o una ruta de evacuación quedar bloqueada.|La coordinación debe revisar todas las dimensiones, no solo la proximidad física.
75|Aclarar funciones de empresa titular, concurrentes y principal sin trasladar responsabilidades indebidamente.|La titular informa sobre el centro; cada concurrente comunica sus riesgos y aplica medidas; la principal cumple su deber de vigilancia cuando corresponde. Coordinar no borra obligaciones propias.|El intercambio de información debe traducirse en controles de campo.
76|Organizar la coordinación en cinco momentos desde contratación hasta entrega final.|Antes de contratar se definen interfaces; antes de entrar se comprueban requisitos; antes del trabajo se revisa realidad; durante se supervisa; al cerrar se entrega en estado seguro.|La coordinación es continua, no una reunión única.
77|Conducir una reunión previa que produzca acuerdos verificables.|Se fijan alcance, límites físicos, responsables, secuencia, energía, tráfico, comunicaciones, emergencias, parada y entrega. Los acuerdos se confirman con planos o referencias claras.|Una reunión eficaz termina con decisiones, responsables y criterios de cambio.
78|Controlar permisos y consignaciones compartidas entre equipos relacionados.|Bloquear solo el motor local no protege frente a alimentación desde aguas arriba, gravedad, material o pruebas remotas. La custodia de bloqueos y el cambio de turno deben quedar definidos.|Nadie retira el bloqueo de otra persona sin aplicar el procedimiento autorizado.
79|Reconocer interferencias frecuentes en exterior antes de que coincidan en campo.|Voladura con topografía, carga con peatones, producción con mantenimiento, excavación con líneas y trabajo nocturno con señalización son combinaciones previsibles.|Lo previsible debe planificarse, no resolverse mediante improvisación radiofónica.
80|Coordinar emergencias entre empresas con alarma, recuento, acceso y mando compatibles.|Todas las personas deben reconocer la alarma, conocer el punto de reunión y aparecer en el recuento. Los servicios externos necesitan coordenadas, acceso y guía.|Una emergencia revela inmediatamente si la coordinación era real o documental.
81|Plantear el caso de mantenimiento y producción en un transportador conectado.|La contrata ha bloqueado un motor, pero producción quiere probar el triturador aguas arriba. El alumno debe identificar alimentación de material, energía compartida y necesidad de consignación conjunta.|La prueba no comienza hasta confirmar que todas las personas están fuera y todos los bloqueos se gestionan.
82|Completar la comprobación del bloque con tres escenarios concretos y resolverlos.|Escenario uno: topografía entra en una zona de carga. Escenario dos: mantenimiento bloquea solo el motor local. Escenario tres: dos empresas usan alarmas distintas en el mismo sector.|La coordinación eficaz modifica el trabajo y deja evidencia de los acuerdos.
83|Ordenar el marco normativo por niveles y traducir referencias en obligaciones verificables.|Ante una decisión sobre una pista, el técnico puede necesitar normativa minera, proyecto, Documento sobre Seguridad y Salud, disposición interna, fabricante y evaluación preventiva.|Conocer la norma significa saber dónde encontrarla y cómo afecta al trabajo.
84|Situar el curso en la ITC 02.1.02 y la Especificación Técnica 2003-1-10.|El itinerario del grupo cinco punto dos, letra a, distribuye veinte horas en seis bloques. Esta trazabilidad permite demostrar que el programa cubre el contenido mínimo.|La presentación y el guion deben conservar la denominación y el reparto oficiales.
85|Explicar presencialidad, reciclaje mínimo y periodicidad para técnicos titulados de exterior.|La formación regulada es presencial; el reciclaje tiene al menos cinco horas y, para este grupo, una frecuencia máxima de cuatro años. Audios y plataforma apoyan la sesión, pero no la sustituyen.|El diseño digital debe integrarse en una impartición presencial documentada.
86|Presentar el Reglamento General de Normas Básicas de Seguridad Minera y las ITC aplicables.|La normativa minera se concreta en proyecto, geometría, organización, polvo y disposiciones internas. Una costumbre local no puede reducir una exigencia vigente.|Siempre se consulta el texto aplicable y las condiciones de la autoridad minera.
87|Explicar derechos, obligaciones y principios preventivos de la Ley 31 de 1995 y su desarrollo.|El trabajador tiene derecho a protección y formación y debe usar correctamente medios y comunicar peligros; la empresa evalúa, planifica e integra la prevención. El servicio preventivo asesora, pero no sustituye la gestión.|La prevención es una responsabilidad organizada y no una decisión aislada del técnico.
88|Relacionar las normas de equipos, EPI y coordinación con decisiones concretas.|El Real Decreto 1215 gobierna adecuación y uso de equipos; el 773 regula EPI; el 171 organiza concurrencia empresarial. Una misma tarea puede activar los tres.|La norma se convierte en selección, comprobación, información y supervisión.
89|Relacionar polvo, sílice, ruido y vibraciones con evaluación, medición y vigilancia de la salud.|Cada agente tiene método, criterio y medidas específicas. No se intercambian límites ni se presume cumplimiento por ausencia de quejas.|La medición representa condiciones reales y conduce a una decisión preventiva.
90|Aclarar que el curso no habilita para explosivos ni para trabajos eléctricos reservados.|El técnico puede coordinar exclusiones y verificar condiciones, pero artilleros y personal eléctrico deben poseer la habilitación o cualificación correspondiente.|Formación preventiva general no amplía competencias legales específicas.
91|Distinguir norma jurídica, guía, nota técnica, fabricante, proyecto e instrucción interna.|Una guía ayuda a interpretar y una nota técnica propone buenas prácticas; la norma obliga en su ámbito; el fabricante fija límites del equipo; proyecto, Documento sobre Seguridad y Salud e instrucciones internas concretan el centro.|Ninguna recomendación se presenta como obligación literal sin comprobar su fuente.
92|Cerrar con un método de revisión normativa y cultura de actualización.|Se registra la fuente oficial, se vigilan cambios, se analiza impacto, se aprueba y comunica la adaptación y se retiran copias obsoletas. Además, se revisan derechos y obligaciones, instrucciones de trabajo y disposiciones internas.|La normativa solo protege cuando la organización la mantiene vigente y la lleva al campo.
"""


def parse_meta():
    out = {}
    for raw in META_RAW.strip().splitlines():
        n, purpose, example, key = raw.split("|", 3)
        out[int(n)] = {"purpose": purpose, "example": example, "key": key}
    return out


STOPWORDS = set(
    "de la el y en a los las del un una que para por con se su al o como es no sobre más esta este "
    "cada lo entre desde cuando qué debe hay son ser sus una unos unas dentro ante hacia tras durante"
    .split()
)


def tokens(text: str):
    return {
        w
        for w in re.findall(r"[a-záéíóúüñ]{4,}", text.lower())
        if w not in STOPWORDS
    }


def clean_text(text: str):
    text = text.replace("", ". ").replace("−", ". ").replace("•", ". ")
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s+([,.;:])", r"\1", text)
    text = re.sub(r"\.{2,}", ".", text)
    return text


def strip_section_lead(text: str):
    text = clean_text(text)
    text = re.sub(r"^(?:\d+(?:\.\d+)?\s+)?[^.]{1,80}(?=\s+[A-ZÁÉÍÓÚÜÑ])", "", text).strip()
    return text


def load_slides():
    raw = SLIDES_TXT.read_text(errors="replace").split("\f")
    slides = []
    for chunk in raw:
        lines = [clean_text(line) for line in chunk.splitlines() if clean_text(line)]
        if lines:
            slides.append({"title": lines[0], "lines": lines[1:]})
    assert len(slides) == 92, len(slides)
    return slides


def load_manual_parts():
    manual = MANUAL_TXT.read_text(errors="replace").replace("\f", "\n\n")
    parts = []
    for raw in re.split(r"\n\s*\n", manual):
        p = clean_text(" ".join(x.strip() for x in raw.splitlines() if x.strip()))
        if len(p.split()) < 25:
            continue
        if "Edición 1.1" in p or "FORMACIÓN PREVENTIVA ·" in p:
            continue
        parts.append(p)
    return parts


def block_for_slide(n):
    if n <= 10:
        return 0
    if n <= 20:
        return 1
    if n <= 42:
        return 2
    if n <= 62:
        return 3
    if n <= 72:
        return 4
    if n <= 82:
        return 5
    return 6


BLOCK_NAMES = {
    0: "Marco de partida",
    1: "Definición de los trabajos",
    2: "Técnicas preventivas y de protección",
    3: "Equipos de trabajo EPI y medios auxiliares",
    4: "Control y vigilancia del entorno",
    5: "Interferencias con otras actividades",
    6: "Normativa y legislación",
}


BLOCK_PARTS = {
    0: list(range(3, 9)),
    1: list(range(9, 13)),
    2: list(range(13, 33)),
    3: list(range(33, 42)),
    4: list(range(42, 46)),
    5: list(range(46, 50)),
    6: list(range(50, 55)),
}


MANUAL_OVERRIDE = {
    1: [3, 5], 2: [4, 5], 3: [0, 1], 4: [2, 12], 5: [6], 6: [7], 7: [8], 8: [8],
    9: [4, 6, 7], 10: [4, 6, 7],
    11: [9, 11], 12: [11, 12], 13: [12], 14: [12], 15: [12], 16: [12, 48],
    17: [12, 45], 18: [12], 19: [9, 10, 11, 12], 20: [9, 10, 11, 12],
    21: [15], 22: [16, 17], 23: [16, 17], 24: [16, 17], 25: [18, 19], 26: [18, 19],
    27: [20, 21], 28: [20, 21], 29: [20, 21], 30: [20, 21], 31: [22], 32: [22],
    33: [23], 34: [23], 35: [23], 36: [24, 25], 37: [24, 25], 38: [26, 27],
    39: [28, 29, 30, 31], 40: [32], 41: [13, 14, 23, 24, 28, 32], 42: [13, 14, 23, 24, 28, 32],
    43: [33, 34], 44: [33, 34], 45: [33, 34], 46: [33, 34], 47: [35],
    48: [33, 38], 49: [33, 36], 50: [33], 51: [36],
    52: [38], 53: [36, 38], 54: [39], 55: [39], 56: [40], 57: [40],
    58: [], 59: [], 60: [41], 61: list(range(34, 42)), 62: list(range(34, 42)),
    63: [42, 43], 64: [42, 43], 65: [44], 66: [44, 45], 67: [44, 45],
    68: [45], 69: [45], 70: [44, 45], 71: list(range(42, 46)), 72: list(range(42, 46)),
    73: [46, 48], 74: [46, 48], 75: [47, 48], 76: [47, 48], 77: [47, 48], 78: [47, 48],
    79: [49], 80: [46, 47, 48, 49], 81: [49], 82: [49],
    83: [50, 51], 84: [52], 85: [52], 86: [50, 51], 87: [50, 51],
    88: [50, 51], 89: [50, 51], 90: [50, 51], 91: [50, 51, 54], 92: [53, 54],
}


CORE_OVERRIDES = {
    19: (
        "Primera pregunta. Ordena las fases preparación, arranque, carga, transporte y tratamiento. "
        "Segunda pregunta. Explica qué falta en una ficha que no incluye criterios de parada ni autoridad de reanudación. "
        "Tercera pregunta. Decide si puede iniciarse una tarea cuando no se conoce quién autoriza la reapertura. "
        "Cuarta pregunta. Enumera las fuentes de energía que deben revisarse: mecánica, eléctrica, hidráulica, neumática, térmica, química y gravitatoria, además de cualquier otra presente en la tarea."
    ),
    61: (
        "Primera pregunta. Clasifica como crítico, limitante o menor un fallo de freno, un daño meramente cosmético, "
        "un cinturón defectuoso o no utilizado en un equipo con riesgo de vuelco y un elemento secundario cuya función se haya confirmado como no necesaria para la condición autorizada. "
        "Segunda pregunta. Decide si el marcado CE garantiza el uso seguro en cualquier terreno. "
        "Tercera pregunta. Explica con qué se asegura un implemento elevado. "
        "Cuarta pregunta. Decide qué hacer cuando una trituradora dispara repetidamente una protección y se rearma sin investigar la causa."
    ),
    71: (
        "Primera pregunta, de elección múltiple. Ante un nivel de alarma, elige entre: continuar vigilando; limitarse a registrar; "
        "parar, aislar, excluir, evaluar y autorizar la reapertura; o esperar al final del turno. "
        "Segunda pregunta. Ordena estas fases: preparar el estándar, observar desde zona segura, clasificar la desviación y verificar el cierre. "
        "Tercera pregunta. Producción presiona para reanudar una zona cerrada. Indica quién puede autorizarlo y con qué evidencia. "
        "Cuarta pregunta. Razona si una fotografía basta por sí sola para cerrar una desviación."
    ),
    82: (
        "Comprobación mediante tres escenarios. Uno: topografía entra en una zona de carga sin contacto positivo. "
        "Dos: mantenimiento bloquea solo el motor local de un transportador mientras producción quiere probar el equipo aguas arriba. "
        "Tres: dos empresas usan alarmas y puntos de reunión diferentes en el mismo sector. "
        "En el primer escenario se separan actividades y se confirma el acceso con el operador. En el segundo se establece consignación conjunta y procedimiento de prueba. "
        "En el tercero se unifican alarma, evacuación, recuento y mando. La coordinación eficaz controla interferencias; no se limita a intercambiar documentos."
    ),
}


# Párrafos de precisión para contenidos que deben pronunciarse completos aunque
# la diapositiva original los presente en forma de tabla o esquema.
EXTRA_PARAGRAPHS = {
    27: [
        "En este punto conviene fijar las magnitudes reglamentarias. En el desarrollo normal del trabajo, la plataforma debe permitir las maniobras manteniendo una distancia mínima de cinco metros al borde del banco. El arcén entre el borde de la pista o del acceso y el pie o borde inferior de un talud no puede ser menor de dos metros. En una pista de un solo carril, la calzada tendrá como mínimo una vez y media la anchura del vehículo mayor previsto; con tráfico intenso y pesado, dos veces esa anchura; y, en una pista de dos carriles, tres veces la del vehículo más ancho. Las pendientes longitudinales medias de las pistas no deben superar el diez por ciento, con máximos puntuales del quince por ciento. En accesos a tajos u otros casos especiales puede superarse ese límite si se justifican las condiciones reales, pero nunca el veinte por ciento y siempre con medidas específicas de seguridad. Estos valores son mínimos reglamentarios: el proyecto aprobado y las condiciones reales pueden exigir más protección."
    ],
    39: [
        "Para un trabajo sin tensión, una vez identificados la zona y los elementos de la instalación, la secuencia reglamentaria contiene cinco etapas: desconectar; prevenir cualquier posible realimentación; verificar la ausencia de tensión; poner a tierra y en cortocircuito; y proteger frente a elementos próximos en tensión, cuando existan, estableciendo además la señalización que delimita la zona de trabajo. Hasta completar las cinco etapas no se autoriza el inicio y la parte afectada se considera en tensión. Las maniobras corresponden a personal autorizado y, en alta tensión, a personal cualificado."
    ],
    83: [
        "Ordenemos el mapa normativo. En el nivel minero aparecen la Ley de Minas, el Reglamento General de Normas Básicas de Seguridad Minera, las instrucciones técnicas del capítulo siete, la instrucción cero dos punto uno punto cero dos, la especificación técnica dos mil tres guión uno guión diez y la instrucción cero dos punto cero punto cero dos. El marco preventivo general lo forman la Ley treinta y uno de mil novecientos noventa y cinco y el Real Decreto treinta y nueve de mil novecientos noventa y siete. El Real Decreto mil trescientos ochenta y nueve de mil novecientos noventa y siete concreta las disposiciones mínimas para actividades extractivas. A partir de ahí se añaden las normas específicas de equipos, protección individual, coordinación, agentes, explosivos y electricidad. No se trata de memorizar una lista: se trata de saber qué norma consultar para cada decisión."
    ],
    84: [
        "Para el grupo cinco punto dos, apartado a, la especificación técnica dos mil tres guión uno guión diez fija veinte horas de formación inicial. La distribución oficial es una hora para definir los trabajos; siete horas para técnicas preventivas y de protección; siete horas para equipos de trabajo, protección individual y medios auxiliares; dos horas para control y vigilancia del entorno; dos horas para interferencias con otras actividades; y una hora para normativa y legislación. Esta arquitectura es la que ordena el curso completo."
    ],
    85: [
        "Hay tres datos que no deben confundirse. La formación regulada por la instrucción cero dos punto uno punto cero dos es presencial. El reciclaje tiene una duración mínima de cinco horas lectivas. Y, para los técnicos titulados de exterior del grupo cinco punto dos, apartado a, la periodicidad máxima establecida es de cuatro años. Por tanto, estos audios apoyan y enriquecen la sesión, pero no sustituyen la asistencia, las actividades dirigidas ni la comprobación del aprendizaje."
    ],
    86: [
        "En el nivel específicamente minero, el Reglamento General de Normas Básicas de Seguridad Minera, aprobado por el Real Decreto ochocientos sesenta y tres de mil novecientos ochenta y cinco, proporciona el marco. Las instrucciones técnicas del capítulo siete desarrollan reglas para trabajos a cielo abierto; la instrucción cero dos punto cero punto cero dos aborda la protección frente al polvo; y la instrucción cero dos punto uno punto cero dos regula la formación preventiva. Antes de aplicar una cifra o una obligación hay que revisar el texto consolidado, el proyecto aprobado y las condiciones impuestas por la autoridad minera."
    ],
    88: [
        "La correspondencia operativa es directa. El Real Decreto mil doscientos quince de mil novecientos noventa y siete se aplica a la adecuación, utilización y comprobación de los equipos de trabajo. El Real Decreto setecientos setenta y tres de mil novecientos noventa y siete regula la selección, el uso y el mantenimiento de los equipos de protección individual. Y el Real Decreto ciento setenta y uno de dos mil cuatro desarrolla la coordinación de actividades empresariales cuando concurren empresas o trabajadores autónomos. Cada una cubre una capa diferente y ninguna sustituye a las otras."
    ],
    89: [
        "Asocia cada agente con su referencia. La exposición a agentes cancerígenos, incluida la sílice cristalina respirable generada en un procedimiento de trabajo, se encuadra en el Real Decreto seiscientos sesenta y cinco de mil novecientos noventa y siete. El ruido se regula por el Real Decreto doscientos ochenta y seis de dos mil seis. Las vibraciones mecánicas, por el Real Decreto mil trescientos once de dos mil cinco. Y la protección frente al polvo en la explotación minera se desarrolla además en la instrucción técnica complementaria cero dos punto cero punto cero dos. La medición, el control de la fuente y la vigilancia de la salud forman un sistema, no actuaciones aisladas."
    ],
    90: [
        "En explosivos, la referencia general es el Reglamento de Explosivos aprobado por el Real Decreto ciento treinta de dos mil diecisiete. En riesgo eléctrico, la referencia preventiva es el Real Decreto seiscientos catorce de dos mil uno. Este curso no habilita para actuar como artillero, manipular explosivos ni realizar trabajos eléctricos reservados a personal autorizado o cualificado. La función del técnico es reconocer la interferencia, mantener distancias y exclusiones, coordinar con las personas habilitadas y detener la tarea cuando no puede demostrarse una condición segura."
    ],
    91: [
        "Cuando varias fuentes regulan una misma tarea, no debe aplicarse una jerarquía simplificada como si una referencia anulara automáticamente a las demás. Se cumplen la legislación vigente y las condiciones de la autoridad minera; el proyecto, el Documento sobre Seguridad y Salud y las disposiciones internas concretan cómo se trabaja en ese centro; y el manual del fabricante fija límites esenciales del equipo. Ningún documento interno puede reducir una exigencia legal ni autorizar un uso contrario al fabricante. Si aparece una contradicción, se conserva el estado seguro, se detiene la decisión y se eleva a la competencia técnica o jurídica correspondiente."
    ],
}


def relevant_manual(slide, n, parts, max_words=245):
    candidates = MANUAL_OVERRIDE.get(n, BLOCK_PARTS[block_for_slide(n)])
    query = tokens(slide["title"] + " " + " ".join(slide["lines"][:8]))
    scored = []
    for j in candidates:
        overlap = len(query & tokens(parts[j]))
        scored.append((overlap, -abs(j - candidates[0]), j))
    scored.sort(reverse=True)
    selected = []
    words = 0
    for _, _, j in scored:
        para = strip_section_lead(parts[j])
        para = re.sub(r"\s+(?:Fuente|Referencia):.*$", "", para, flags=re.I)
        para = para.lstrip(". ;:")
        count = len(para.split())
        if words and words + count > max_words:
            continue
        selected.append(para)
        words += count
        if words >= 150 or len(selected) >= 3:
            break
    return selected


def slide_prose(slide, max_words=165, n=None):
    if n in CORE_OVERRIDES:
        return CORE_OVERRIDES[n]
    kept = []
    words = 0
    for line in slide["lines"]:
        # Descarta encabezados de tabla aislados y referencias editoriales.
        if line.lower().startswith("referencia:"):
            continue
        if len(line.split()) <= 2 and line.upper() == line:
            continue
        line = re.sub(r"^(\d{1,2})\s+", "", line)
        if not re.search(r"[a-záéíóúüñ]", line.lower()):
            continue
        if words + len(line.split()) > max_words:
            break
        kept.append(line)
        words += len(line.split())
    text = clean_text(" ".join(kept))
    text = re.sub(r"\bDSS\b", "Documento sobre Seguridad y Salud", text)
    text = re.sub(r"\bDIS\b", "disposición interna de seguridad", text)
    text = re.sub(r"\bEPI\b", "equipo de protección individual", text)
    return text


OPENERS = [
    "Vamos a detenernos en una idea que condiciona muchas decisiones de campo.",
    "Esta diapositiva propone mirar el trabajo con criterio técnico y no por costumbre.",
    "El contenido que aparece en pantalla debe convertirse en una forma concreta de decidir.",
    "Antes de pasar a la siguiente operación, conviene entender bien este punto.",
    "Aquí abordamos una situación cotidiana que puede tener consecuencias importantes.",
    "La clave de esta diapositiva está en pasar de la descripción a la comprobación.",
    "En una explotación exterior, este asunto rara vez aparece de forma aislada.",
    "Pensemos en cómo se aplicaría este criterio durante un turno real.",
]


BRIDGES = [
    "Llevado al terreno, esto significa lo siguiente.",
    "Para aplicarlo de forma consistente, conviene ordenar el razonamiento.",
    "El criterio técnico que hay detrás puede resumirse así.",
    "La diapositiva recoge varios elementos que deben leerse como un sistema.",
    "No se trata de memorizar una lista, sino de comprender la relación entre sus partes.",
    "La comprobación práctica comienza observando qué puede cambiar y a quién puede afectar.",
]


def depth_paragraphs(n, m):
    """Añade profundidad oral con variación suficiente entre diapositivas."""
    purpose = m["purpose"]
    key = m["key"]
    variants = [
        (
            "Para trabajar este contenido con método, ordénalo en tres momentos. Antes de empezar, confirma el estándar, las personas afectadas y las condiciones previas. Durante la tarea, observa si aparecen cambios, si las barreras siguen presentes y si alguien entra en la trayectoria del peligro. Al terminar, comprueba el estado de entrega, comunica defectos pendientes y conserva la evidencia necesaria. Esta secuencia evita que la prevención se limite al instante inicial y ayuda a detectar desviaciones que nacen cuando la operación ya está en marcha.",
            f"Un error frecuente sería dar por cumplido el objetivo solo porque existe una instrucción. En este caso, el objetivo real es {purpose[0].lower() + purpose[1:]} La supervisión debe confirmar que el documento puede aplicarse en las condiciones actuales, que los medios están disponibles y que las personas comprenden qué les obliga a detenerse. Por eso, {key[0].lower() + key[1:]}"
        ),
        (
            "Desde la función técnica conviene separar tres preguntas. Primera: cuál es la condición aceptable y en qué documento está definida. Segunda: qué señal mostraría que esa condición se está perdiendo. Tercera: qué acción inmediata impide que la desviación alcance a una persona. Si una de estas preguntas no tiene respuesta, todavía falta preparación. La calidad del control depende de que la respuesta pueda comunicarse con claridad entre turnos, empresas y responsables.",
            f"También hay que evitar una trampa habitual: confundir actividad con resultado. Revisar, informar o señalizar son acciones, pero solo son eficaces si producen una condición controlada. Aplicado a esta diapositiva, necesitamos {purpose[0].lower() + purpose[1:]} La evidencia debe mostrar el resultado conseguido, no únicamente que alguien realizó una gestión. El criterio final es este: {key}"
        ),
        (
            "Piensa ahora en la cadena de barreras. Una primera barrera intenta evitar el suceso; otra limita la exposición; una tercera permite detectar el deterioro; y la respuesta de parada impide continuar cuando las anteriores no ofrecen confianza. No todas las tareas necesitarán cuatro barreras diferentes, pero el razonamiento ayuda a descubrir dependencias excesivas. Si todo descansa en que una persona recuerde una orden en el momento exacto, el sistema es frágil y debe reforzarse desde el diseño o la organización.",
            f"Para supervisar correctamente, no preguntes solo si se ha entendido la explicación. Pide que la persona señale el peligro, describa la trayectoria, identifique el control crítico y explique cómo sabrá que ha fallado. Esa conversación convierte {purpose.lower()} en una competencia observable. La idea que debe poder repetirse sin mirar la pantalla es: {key}"
        ),
        (
            "Este criterio debe mantenerse cuando cambian la producción, el personal o el entorno. Por eso es útil definir qué datos se revisan al inicio de turno, qué observaciones se repiten durante la operación y qué eventos activan una revisión extraordinaria. La lluvia, una avería, una modificación, un relevo de personal o una tarea simultánea pueden invalidar una condición que era aceptable unas horas antes. La evaluación no desaparece al firmarse; acompaña al trabajo mientras exista exposición.",
            f"El técnico aporta valor cuando convierte {purpose.lower()} en decisiones que los demás pueden ejecutar. Eso exige instrucciones breves, responsabilidades claras, comunicación cerrada y una condición de parada que no dependa de negociar bajo presión. Si aparece incertidumbre, se conserva la situación segura y se escala. La razón es sencilla: {key}"
        ),
        (
            "Observa además la diferencia entre presencia y eficacia. Una barrera puede estar físicamente instalada y no cumplir su función: un pretil mal dimensionado, una cámara sucia, una guarda puenteada, un filtro saturado o una señal que nadie ve. Verificar significa comparar con un criterio, probar cuando corresponda y registrar la desviación. Esta mirada funcional evita aceptar como seguro algo que únicamente parece completo desde lejos.",
            f"En la práctica, {purpose.lower()} requiere combinar conversación con las personas, observación del trabajo y consulta de la documentación aplicable. Ninguna de esas fuentes basta por sí sola. Si lo escrito no puede ejecutarse, se revisa; si la práctica contradice el estándar, se controla y se corrige; si el estándar es insuficiente, se gestiona el cambio. De nuevo, el mensaje central es: {key}"
        ),
        (
            "Conviene pensar también en el relevo y la continuidad. La persona que recibe la tarea debe conocer qué se ha hecho, qué controles están activos, qué defectos permanecen abiertos y qué acciones no pueden modificarse sin autorización. Una comunicación incompleta puede reintroducir energía, abrir una zona excluida o eliminar una restricción todavía necesaria. Por eso la información crítica se confirma y se registra en el soporte previsto por la explotación.",
            f"El razonamiento no termina al elegir una medida. Para {purpose.lower()}, asigna siempre quién la implanta, quién la comprueba, durante cuánto tiempo debe mantenerse y qué evidencia autoriza su retirada. Esta disciplina evita controles huérfanos y reaperturas por costumbre. El principio que orienta todas esas decisiones es: {key}"
        ),
    ]
    return list(variants[(n - 1) % len(variants)])


def transfer_paragraphs(n, m):
    variants = [
        (
            "Modifica mentalmente una condición del ejemplo. Imagina que cambia el turno, comienza a llover, aparece una contrata o falla una ayuda técnica. ¿Seguiría siendo válido el mismo control? Si la respuesta depende de la nueva condición, esa dependencia debe figurar en el procedimiento o en el criterio de parada. Esta pregunta prepara al equipo para reconocer el momento en que una tarea deja de ser rutinaria.",
            "Para cerrar el razonamiento, identifica una evidencia previa, una evidencia durante el trabajo y una evidencia final. Pueden ser una inspección registrada, una medición, una comunicación confirmada, una prueba funcional o la aceptación documentada del estado de entrega. Elegirlas antes de empezar evita buscar justificaciones después y permite que otra persona audite la decisión."
        ),
        (
            "Plantea ahora el contenido desde el punto de vista de tres personas: quien ejecuta, quien supervisa y quien puede quedar afectado sin participar en la tarea. La primera necesita instrucciones y medios; la segunda, criterios y autoridad; la tercera, separación e información. Cuando solo se considera al ejecutante, suelen quedar fuera peatones, conductores, contratas, visitantes o trabajadores de otra instalación.",
            "La pregunta útil para el final del turno es: ¿qué condición podría encontrar mañana una persona que no estuvo hoy? La respuesta determina qué debe señalizarse, bloquearse, comunicarse o registrarse. Un trabajo no queda cerrado si transmite a la siguiente persona un riesgo que no puede reconocer o controlar."
        ),
        (
            "Imagina que alguien propone continuar porque la tarea durará solo unos minutos. La corta duración no elimina energía, gravedad, tráfico ni posibilidad de fallo. Puede reducir exposición, pero también favorece la improvisación. La decisión correcta compara el riesgo real con los controles disponibles y mantiene el mismo criterio de parada, aunque la intervención parezca breve.",
            "Comprueba además la calidad de la comunicación. Una orden como vigila la zona o hazlo con cuidado no define límites. Una comunicación operativa indica quién, dónde, desde cuándo, hasta qué condición, por qué canal y qué respuesta confirma el acuerdo. Esa precisión es especialmente importante con ruido, cambios de turno o empresas concurrentes."
        ),
        (
            "Cambia ahora la pregunta de podemos hacerlo a qué tendría que ser cierto para hacerlo con seguridad. Esa formulación obliga a enumerar terreno estable, equipo adecuado, personal competente, zona controlada, energía aislada, comunicación disponible y autorización cuando corresponda. Si no puede comprobarse una condición necesaria, la tarea permanece detenida o se rediseña.",
            "Después de implantar el control, busca indicios de degradación. Las barreras físicas se dañan, los drenajes se obstruyen, las marcas se borran, las alarmas se normalizan y los procedimientos pueden quedarse atrás respecto al trabajo real. La supervisión útil no solo verifica presencia; comprueba que la barrera sigue siendo capaz de cumplir su función."
        ),
        (
            "Supón que el resultado ha sido correcto varias veces. Esa experiencia puede aportar información, pero no demuestra que el riesgo esté controlado. Un evento de baja frecuencia puede conservar consecuencias graves y una desviación puede no haber coincidido todavía con una persona expuesta. La ausencia de accidente es un dato, no una prueba suficiente de seguridad.",
            "Por eso conviene pedir una explicación en sentido inverso: partiendo del daño posible, reconstruye la trayectoria y localiza qué barrera lo impide. Si no aparece una barrera concreta, verificable y mantenida, la conclusión optimista debe revisarse. Este ejercicio ayuda a detectar controles que existen solo en el lenguaje."
        ),
        (
            "Traslada este criterio a una tarea no rutinaria. En una avería, una limpieza extraordinaria o una modificación, las referencias habituales pueden no describir exactamente la situación. El equipo debe detenerse, delimitar el alcance, identificar energías y consultar a quien tenga competencia para aprobar el método. La urgencia no crea una excepción preventiva.",
            "Finalmente, piensa cómo formarías a una persona nueva sobre este punto. No bastaría con mostrarle la diapositiva. Tendría que reconocer el escenario en campo, explicar cuándo parar, localizar el documento aplicable y demostrar el control bajo supervisión. Esa es la diferencia entre haber recibido información y haber adquirido capacidad para trabajar con seguridad."
        ),
    ]
    return list(variants[(n - 1) % len(variants)])


def case_walkthrough_paragraphs(m):
    purpose = m["purpose"]
    key = m["key"]
    return [
        f"Descompongamos el ejemplo paso a paso. Primero se fija la situación inicial y se evita añadir datos que no conocemos. Después se localiza la fuente de daño y se dibuja mentalmente su trayectoria. A continuación se identifican las personas, equipos o instalaciones que podrían quedar dentro. Solo entonces se revisan las barreras existentes y se decide si son suficientes. Este orden es importante para {purpose[0].lower() + purpose[1:]} Si empezamos proponiendo una solución, podemos pasar por alto una energía o una exposición que cambia por completo la decisión.",
        "En el siguiente paso se elige una medida inmediata y otra de estabilización. La medida inmediata puede consistir en parar, retirar personas, inmovilizar, aislar o delimitar. La estabilización aborda la causa mediante reparación, rediseño, drenaje, mantenimiento, nueva secuencia o coordinación. Ambas necesitan responsable y plazo. Finalmente se define una prueba de eficacia: observar bajo condiciones representativas, medir, realizar una prueba funcional o confirmar el estado mediante personal competente.",
        f"Si esta situación surgiera durante tu turno, explica en voz alta qué comunicarías. Una comunicación completa podría comenzar indicando el lugar exacto, la condición observada y la zona potencialmente afectada. Continuaría con la medida adoptada, las personas avisadas y lo que debe ocurrir antes de reanudar. Esa estructura reduce malentendidos y deja claro que {key[0].lower() + key[1:]}"
    ]


def activity_depth(m):
    return [
        "Durante el trabajo en grupo, asigna una persona para describir los hechos, otra para cuestionar las suposiciones y otra para comprobar la trazabilidad con documentos del centro. Si trabajas individualmente, realiza esas tres pasadas por separado. En la primera no propongas soluciones; en la segunda identifica incertidumbres; en la tercera elige controles siguiendo la jerarquía y define cómo se verificarán. Esta separación ayuda a que una solución intuitiva no oculte datos que todavía faltan.",
        f"La respuesta final debe poder comunicarse por radio o en una reunión previa sin ambigüedad. Incluye ubicación, condición detectada, zona afectada, medida inmediata, persona responsable y requisito de reapertura. Después contrástala con este principio: {m['key']} Si tu propuesta depende de que todos recuerden tener cuidado, busca una barrera técnica u organizativa más robusta."
    ]


def check_depth(m):
    return [
        "Para cada respuesta, escribe también cuál sería el error más peligroso y qué consecuencia podría producir. Este paso evita acertar por descarte sin comprender el control. Cuando una pregunta admita condiciones, indícalas: muchas medidas solo son válidas si el terreno, el equipo, la competencia y la autorización coinciden con lo previsto.",
        f"Al terminar, selecciona la pregunta que más dudas te haya generado y vuelve al documento aplicable. La formación debe enseñarte a localizar y verificar, no a improvisar una certeza. Mantén como referencia esta idea: {m['key']}"
    ]


def solution_depth(m):
    return [
        "Una resolución sólida distingue control inmediato y solución definitiva. El control inmediato detiene la exposición, pero puede dejar la causa intacta. La solución definitiva modifica diseño, método, mantenimiento u organización y necesita una comprobación posterior. No retires la restricción temporal hasta que esa comprobación exista y la persona autorizada confirme la reapertura.",
        f"Por último, traslada el aprendizaje a otra tarea. Pregunta qué escenario diferente comparte la misma energía, trayectoria o fallo de coordinación. Esa transferencia es lo que convierte una respuesta de curso en capacidad profesional. El criterio que debe conservarse es: {m['key']}"
    ]


def is_activity(title):
    return "ACTIVIDAD DEL ALUMNO" in title or title.startswith("ACTIVIDAD ")


def is_check(title):
    return "COMPROBACIÓN" in title and "RESPUESTAS" not in title and "RESOLUCIÓN" not in title


def is_solution(title):
    return "RESPUESTAS" in title or "RESOLUCIÓN" in title or "RESOLUCIONES" in title


def activity_minutes(title):
    m = re.search(r"(\d+)\s+MINUTOS", title)
    return int(m.group(1)) if m else None


def clip_sentences(text, max_words):
    """Conserva frases completas siempre que sea posible y respeta un techo de palabras."""
    text = clean_text(text)
    if len(text.split()) <= max_words:
        return text
    sentences = re.split(r"(?<=[.!?])\s+", text)
    selected = []
    words = 0
    for sentence in sentences:
        count = len(sentence.split())
        if selected and words + count > max_words:
            break
        if not selected and count > max_words:
            return " ".join(sentence.split()[:max_words]).rstrip(",;:") + "."
        selected.append(sentence)
        words += count
    return " ".join(selected).strip()


def finalize_short_script(paragraphs, m, min_words=260, max_words=370):
    paragraphs = [clean_text(p) for p in paragraphs if clean_text(p)]
    total = sum(len(p.split()) for p in paragraphs)

    if total < min_words:
        paragraphs.insert(
            -1,
            "Para trasladarlo al puesto, comprueba tres cosas: qué condición debe existir antes de empezar, "
            "qué señal obliga a detenerse y qué evidencia permite cerrar la tarea. Asigna además una persona "
            "responsable de verificar el control. Así, el criterio deja de ser una recomendación general y se "
            "convierte en una decisión que puede observarse, comunicarse y corregirse durante el turno.",
        )
        total = sum(len(p.split()) for p in paragraphs)

    if total < min_words:
        paragraphs.insert(
            -1,
            "Antes de cerrar, explica el criterio con tus propias palabras y aplícalo a una tarea conocida. "
            "Si no puedes identificar la barrera, la persona responsable y la comprobación final, todavía falta "
            "concretar la decisión preventiva.",
        )
        total = sum(len(p.split()) for p in paragraphs)

    # Ajuste final para que ningún audio supere aproximadamente tres minutos.
    while total > max_words:
        candidates = [(len(p.split()), i) for i, p in enumerate(paragraphs[1:-1], 1)]
        if not candidates:
            break
        size, idx = max(candidates)
        reduction = min(size - 35, total - max_words)
        if reduction <= 0:
            break
        paragraphs[idx] = clip_sentences(paragraphs[idx], size - reduction)
        new_total = sum(len(p.split()) for p in paragraphs)
        if new_total >= total:
            break
        total = new_total

    return paragraphs


def narration_for(n, slide, meta, manual_parts):
    title = slide["title"]
    m = meta[n]
    paragraphs = []

    if is_activity(title):
        minutes = activity_minutes(title) or 15
        paragraphs.append(
            f"Llegamos a una actividad práctica de aproximadamente {minutes} minutos. "
            f"El objetivo es {m['purpose'][0].lower() + m['purpose'][1:]}"
        )
        core = clip_sentences(slide_prose(slide, 140, n), 85)
        if core:
            paragraphs.append("La situación de trabajo es esta. " + core)
        paragraphs.append("Para orientar el análisis, considera este encuadre. " + clip_sentences(m["example"], 55))
        paragraphs.append(
            "Trabaja en este orden. Primero separa hechos de suposiciones. Después identifica la energía o condición peligrosa, su trayectoria y las personas expuestas. A continuación propone una medida inmediata para retirar la exposición y otra que corrija la causa. Termina indicando quién decide, qué debe comprobarse y qué condición permite reabrir. Evita respuestas genéricas como tener cuidado o señalizar sin explicar qué barrera controla el escenario."
        )
        paragraphs.append(
            "Detén ahora la reproducción y realiza la actividad. Conserva tu respuesta para compararla con la resolución posterior. " + m["key"]
        )
        return finalize_short_script(paragraphs, m)

    if is_check(title):
        paragraphs.append(
            "Vamos a comprobar el aprendizaje del bloque. Justifica cada respuesta relacionando peligro, exposición, control y evidencia de campo."
        )
        core = clip_sentences(slide_prose(slide, 230, n), 140)
        if core:
            paragraphs.append("Estas son las cuestiones que debes resolver. " + core)
        paragraphs.append("Utiliza este ejemplo para comprobar tu criterio. " + clip_sentences(m["example"], 50))
        paragraphs.append(
            "Haz una pausa después de cada pregunta. Si dos opciones parecen posibles, elige la que elimine antes la exposición, dependa menos de la conducta y permita verificar su eficacia. En una secuencia, conserva el orden del trabajo seguro. Anota también el error más peligroso y la consecuencia que produciría; así evitarás acertar simplemente por descarte."
        )
        paragraphs.append(
            "Detén la reproducción, responde por escrito y explica por qué descartas las demás opciones. Después continúa con la resolución razonada. " + m["key"]
        )
        return finalize_short_script(paragraphs, m)

    if is_solution(title):
        paragraphs.append(
            "Vamos a revisar la resolución. Lo importante es comprobar que la decisión retira la exposición inmediata, trata la causa y define cómo se recupera una condición segura."
        )
        core = clip_sentences(slide_prose(slide, 180, n), 95)
        if core:
            paragraphs.append("La respuesta propuesta se desarrolla así. " + core)
        manual = " ".join(relevant_manual(slide, n, manual_parts, 120))
        if manual:
            paragraphs.append(clip_sentences(manual, 70))
        paragraphs.append("En una situación reconocible ocurriría lo siguiente. " + clip_sentences(m["example"], 55))
        paragraphs.append(
            "Distingue siempre entre control inmediato y solución definitiva. El primero detiene la exposición; la segunda corrige el diseño, el método, el mantenimiento o la organización. No retires una restricción temporal hasta comprobar el resultado y obtener la autorización prevista. Compara ahora esta solución con la tuya: revisa responsable, criterio de parada, condición de reapertura y evidencia de eficacia. " + m["key"]
        )
        return finalize_short_script(paragraphs, m)

    purpose_spoken = m["purpose"][0].lower() + m["purpose"][1:]
    paragraphs.append(OPENERS[(n - 1) % len(OPENERS)] + " Nuestro objetivo es " + purpose_spoken)
    extras = EXTRA_PARAGRAPHS.get(n, [])
    if extras:
        paragraphs.extend(clip_sentences(p, 145) for p in extras)
    else:
        manual = " ".join(relevant_manual(slide, n, manual_parts, 150))
        if manual:
            paragraphs.append(clip_sentences(manual, 105))
    paragraphs.append("Pensemos en un ejemplo. " + clip_sentences(m["example"], 58))
    paragraphs.append(
        "La actuación técnica debe seguir una secuencia sencilla: observar la condición real, retirar o evitar la exposición, aplicar la barrera más robusta disponible y comprobar que funciona. La medida necesita un responsable y un criterio visible de aceptación. Si cambia el terreno, el equipo, el personal o la tarea simultánea, vuelve a evaluar antes de continuar. Una instrucción como hacerlo con cuidado no define un control verificable."
    )
    paragraphs.append(
        "Antes de continuar, pregúntate qué observarías o medirías para demostrar que el criterio se cumple durante el turno. Puede ser una condición visible, una medición, una prueba funcional o un registro con responsable. Resume la diapositiva con esta idea: " + m["key"]
    )
    return finalize_short_script(paragraphs, m)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    normal.font.color.rgb = RGBColor(30, 30, 30)

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    for name, size in (("Heading 1", 18), ("Heading 2", 13)):
        s = styles[name]
        s.font.name = "Aptos Display"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(7)

    if "Locucion" not in styles:
        loc = styles.add_style("Locucion", WD_STYLE_TYPE.PARAGRAPH)
    else:
        loc = styles["Locucion"]
    loc.base_style = styles["Normal"]
    loc.font.name = "Aptos"
    loc._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    loc._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    loc.font.size = Pt(11.25)
    loc.paragraph_format.space_after = Pt(8)
    loc.paragraph_format.line_spacing = 1.14

    if "Meta" not in styles:
        meta = styles.add_style("Meta", WD_STYLE_TYPE.PARAGRAPH)
    else:
        meta = styles["Meta"]
    meta.base_style = styles["Normal"]
    meta.font.name = "Aptos"
    meta._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    meta._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    meta.font.size = Pt(9)
    meta.font.color.rgb = RGBColor(90, 90, 90)
    meta.paragraph_format.space_after = Pt(6)


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("InmínerCampus · Guion maestro de locución · Técnicos titulados de exterior")
    set_font(r, size=8, color="777777")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def estimated_audio_minutes(paragraphs):
    words = sum(len(p.split()) for p in paragraphs)
    return max(2, round(words / 132))


def build():
    slides = load_slides()
    manual_parts = load_manual_parts()
    meta = parse_meta()
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    scripts = []
    for n, slide in enumerate(slides, 1):
        paragraphs = narration_for(n, slide, meta, manual_parts)
        scripts.append(paragraphs)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    configure_styles(doc)
    add_footer(section)
    add_page_number(section.footer.add_paragraph())

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(4.8))

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Guion maestro de locución")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Formación preventiva para técnicos titulados en actividades extractivas de exterior")
    set_font(r, size=16, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ITC 02.1.02 · Grupo 5.2.a · ET 2003-1-10 · Curso inicial de 20 horas")
    set_font(r, size=11, color="555555")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Versión 1.1 · Audios de 2 a 3 minutos · Un archivo por diapositiva · Texto preparado para ElevenLabs")
    set_font(r, size=10, color="777777")

    doc.add_page_break()
    current_block = None
    for n, (slide, paragraphs) in enumerate(zip(slides, scripts), 1):
        block = block_for_slide(n)
        if current_block != block:
            current_block = block

        # Heading 1 is intentionally used so Studio can detect one chapter per slide.
        heading = doc.add_heading(f"Diapositiva {n:02d} {slide['title']}", level=1)
        heading.paragraph_format.page_break_before = n != 1
        heading.paragraph_format.keep_with_next = True

        for text in paragraphs:
            p = doc.add_paragraph(style="Locucion")
            p.paragraph_format.widow_control = True
            p.add_run(text)

    doc.core_properties.title = "Guion maestro de locución para técnicos titulados de exterior"
    doc.core_properties.subject = "ITC 02.1.02 grupo 5.2.a ET 2003-1-10"
    doc.core_properties.author = "InmínerCampus"
    doc.core_properties.keywords = "minería, formación preventiva, exterior, guion, ElevenLabs"
    doc.save(OUT)

    # Consola de control editorial.
    counts = [sum(len(p.split()) for p in script) for script in scripts]
    print(OUT)
    print("slides", len(scripts))
    print("words", sum(counts))
    print("min", min(counts), "max", max(counts), "avg", round(sum(counts) / len(counts), 1))
    print("estimated_audio_minutes", sum(max(2, round(x / 132)) for x in counts))
    print("short", [(i + 1, c) for i, c in enumerate(counts) if c < 260])


if __name__ == "__main__":
    build()
